from datetime import datetime, timezone
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook

from provision.config import budget_for, load_runs, normalize_config, resolve_start
from provision.sheet import _sheet_path, load_spreadsheet, normalize_spreadsheet


def _columns():
    return {
        "account_id": "Account ID",
        "account_name": "Account Name",
        "bucket": "Cost Bucket",
        "amount": "Annual Budget",
        "recipients": "Recipients",
        "alerts": [
            {"header": "Alert 80", "percentage": 80, "based_on": "ACTUAL_COST"},
            {"header": "Alert 100", "percentage": 100, "based_on": "ACTUAL_COST"},
        ],
    }


def _sheet(**overrides):
    sheet = {"key": "account", "sheet": None, "columns": _columns()}
    sheet.update(overrides)
    return sheet


class SheetLoaderTest(unittest.TestCase):
    def test_month_start_header_expands_and_derives_uneven_annual_amount(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "budgets.csv"
            headers = ["Account ID"] + [
                f"{year:04d}-{month:02d}"
                for year, month in (
                    [(2026, month) for month in range(4, 13)]
                    + [(2027, month) for month in range(1, 4)]
                )
            ]
            values = ["111111111111"] + [4000] * 9 + [3900] * 3
            path.write_text(
                ",".join(headers) + "\n" + ",".join(map(str, values)) + "\n",
                encoding="utf-8",
            )
            sheet = normalize_spreadsheet(
                {
                    "file": str(path),
                    "key": "account",
                    "only": True,
                    "columns": {
                        "account_id": "Account ID",
                        "months": "2026-04",
                        "alerts": [],
                    },
                }
            )
            rows = load_spreadsheet(path, sheet)

        self.assertEqual(sheet["columns"]["month_start"], "2026-04-01")
        self.assertEqual(sheet["columns"]["months"][-1], "2027-03")
        self.assertEqual(rows["111111111111"]["amount"], 47700)
        self.assertEqual(rows["111111111111"]["monthly_amounts"][-3:], [3900] * 3)

    def test_csv_reads_thresholds_and_recipients(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "budgets.csv"
            path.write_text(
                "Account ID,Account Name,Cost Bucket,Annual Budget,Alert 80,Alert 100,Recipients\n"
                "012345678901,Testing,hosted,1200,Y,N,a@example.com\n"
                "100000000120,Other,hosted,2400,yes,,\n"
                "333333333333,Third,hosted,\"1,200\",Y,true,\"a@example.com, b@example.com\"\n",
                encoding="utf-8",
            )
            rows = load_spreadsheet(path, _sheet())

        self.assertEqual(rows["012345678901"]["amount"], 1200)
        self.assertEqual(rows["012345678901"]["account_name"], "Testing")
        self.assertEqual(
            [alert["percentage"] for alert in rows["012345678901"]["alerts"]],
            [80],
        )
        self.assertEqual(rows["012345678901"]["alerts"][0]["emails"], ["a@example.com"])
        self.assertEqual(rows["012345678901"]["warnings"], [])
        self.assertEqual(rows["100000000120"]["alerts"][0]["emails"], [])
        self.assertIn(
            "Alert 80 is enabled with no recipients",
            rows["100000000120"]["warnings"],
        )
        self.assertEqual(rows["333333333333"]["amount"], 1200)
        self.assertEqual(
            [alert["percentage"] for alert in rows["333333333333"]["alerts"]],
            [80, 100],
        )
        self.assertEqual(
            rows["333333333333"]["alerts"][0]["emails"],
            ["a@example.com", "b@example.com"],
        )

    def test_xlsx_pads_a_numeric_account_id(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "budgets.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Budgets"
            worksheet.append(
                ["Account ID", "Account Name", "Annual Budget", "Alert 80", "Recipients"]
            )
            worksheet.append([12345678901, "Testing", 500, True, None])
            workbook.save(path)
            workbook.close()
            rows = load_spreadsheet(
                path,
                _sheet(
                    sheet="Budgets",
                    columns={
                        "account_id": "Account ID",
                        "account_name": "Account Name",
                        "amount": "Annual Budget",
                        "recipients": "Recipients",
                        "alerts": [
                            {
                                "header": "Alert 80",
                                "percentage": 80,
                                "based_on": "ACTUAL_COST",
                            }
                        ],
                    },
                ),
            )

        self.assertIn("012345678901", rows)
        self.assertNotIn("Testing (012345678901)", rows)
        self.assertEqual(rows["012345678901"]["alerts"][0]["emails"], [])
        self.assertEqual(
            rows["012345678901"]["warnings"],
            ["Alert 80 is enabled with no recipients"],
        )

    def test_blank_folder_is_kept_as_a_bucket_fallback_warning(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "budgets.csv"
            path.write_text(
                "Account ID,Account Name,Folder,Annual Budget\n"
                "111111111111,Alpha,,100\n",
                encoding="utf-8",
            )
            sheet = normalize_spreadsheet(
                {
                    "file": str(path),
                    "key": "account",
                    "only": True,
                    "columns": {
                        "account_id": "Account ID",
                        "account_name": "Account Name",
                        "folder": "Folder",
                        "amount": "Annual Budget",
                        "alerts": [],
                    },
                }
            )
            rows = load_spreadsheet(path, sheet)

        self.assertIsNone(rows["111111111111"]["folder"])
        self.assertIn(
            "Folder is blank; using the cost bucket as the folder",
            rows["111111111111"]["warnings"],
        )

    def test_combined_label_is_not_parsed(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "budgets.csv"
            path.write_text(
                "Account ID,Annual Budget\n"
                "Testing (012345678901),100\n",
                encoding="utf-8",
            )
            rows = load_spreadsheet(
                path,
                _sheet(
                    columns={
                        "account_id": "Account ID",
                        "amount": "Annual Budget",
                        "alerts": [],
                    }
                ),
            )

        self.assertEqual(list(rows), ["Testing (012345678901)"])

    def test_unmapped_duplicate_and_blank_headers_are_ignored(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "budgets.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["Product", "Account ID", "Note", "Note", None, "FY27 TOTAL"])
            worksheet.append(["Platform", 100000000001, "a", "b", "created manually", 47700])
            worksheet.append(["Total", None, None, None, None, 47700])
            workbook.save(path)
            workbook.close()
            skipped = []
            rows = load_spreadsheet(
                path,
                _sheet(
                    columns={
                        "account_id": "Account ID",
                        "amount": "FY27 TOTAL",
                        "alerts": [],
                    }
                ),
                skipped,
            )

        self.assertEqual(list(rows), ["100000000001"])
        self.assertEqual(rows["100000000001"]["amount"], 47700)
        self.assertEqual(skipped, ["skipped spreadsheet row 3: no Account ID"])

    def test_duplicate_mapped_header_names_the_columns(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "budgets.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["Account ID", "FY27 TOTAL", "Account ID"])
            worksheet.append([111111111111, 10, 111111111111])
            workbook.save(path)
            workbook.close()
            with self.assertRaises(ValueError) as caught:
                load_spreadsheet(
                    path,
                    _sheet(
                        columns={
                            "account_id": "Account ID",
                            "amount": "FY27 TOTAL",
                            "alerts": [],
                        }
                    ),
                )

        self.assertIn("more than one Account ID column (A, C)", str(caught.exception))

    def test_duplicate_account_id_is_an_error(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "budgets.csv"
            path.write_text(
                "Account ID,Annual Budget\n111111111111,10\n111111111111,20\n",
                encoding="utf-8",
            )
            with self.assertRaises(ValueError) as caught:
                load_spreadsheet(
                    path,
                    _sheet(
                        columns={
                            "account_id": "Account ID",
                            "amount": "Annual Budget",
                            "alerts": [],
                        }
                    ),
                )

        self.assertIn("duplicate account id 111111111111", str(caught.exception))

    def test_alert_value_must_be_yes_or_no(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "budgets.csv"
            path.write_text(
                "Account ID,Annual Budget,Alert 80\n111111111111,10,maybe\n",
                encoding="utf-8",
            )
            with self.assertRaises(ValueError):
                load_spreadsheet(
                    path,
                    _sheet(
                        columns={
                            "account_id": "Account ID",
                            "amount": "Annual Budget",
                            "alerts": [
                                {
                                    "header": "Alert 80",
                                    "percentage": 80,
                                    "based_on": "ACTUAL_COST",
                                }
                            ],
                        }
                    ),
                )


class SheetPathTest(unittest.TestCase):
    def test_home_and_absolute_paths_stay_outside_the_config_directory(self):
        home = _sheet_path("examples/testbed.yaml", "~/Downloads/budgets.xlsx", None)
        absolute = _sheet_path("examples/testbed.yaml", "/tmp/budgets.xlsx", None)
        self.assertEqual(home, Path("~/Downloads/budgets.xlsx").expanduser())
        self.assertNotIn("examples", home.parts)
        self.assertEqual(absolute, Path("/tmp/budgets.xlsx"))

    def test_cli_override_replaces_a_missing_configured_file(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            workbook = root / "override.csv"
            workbook.write_text(
                "Account ID,Annual Budget\n999999999999,50\n",
                encoding="utf-8",
            )
            config_path = root / "run.yaml"
            config_path.write_text(
                "cost_category: Example\n"
                "spreadsheet:\n"
                "  file: ~/missing.xlsx\n"
                "  key: account\n"
                "  only: true\n"
                "  columns:\n"
                "    account_id: Account ID\n"
                "    amount: Annual Budget\n"
                "    alerts: []\n",
                encoding="utf-8",
            )
            runs = load_runs(str(config_path), str(workbook))

        self.assertEqual(runs[0]["spreadsheet"]["file"], str(workbook))
        self.assertEqual(set(runs[0]["spreadsheet"]["rows"]), {"999999999999"})

    def test_cli_override_accepts_a_home_path(self):
        resolved = _sheet_path(
            "examples/testbed.yaml",
            "missing.csv",
            "~/Downloads/budgets.xlsx",
        )
        self.assertEqual(resolved, Path("~/Downloads/budgets.xlsx").expanduser())


class SheetBudgetTest(unittest.TestCase):
    def test_month_start_sets_the_annual_period_and_fiscal_start(self):
        config = normalize_config(
            {
                "cost_category": "Example",
                "budget": {
                    "type": "SPECIFIED_AMOUNT",
                    "period": "MONTHLY",
                    "amount": None,
                    "monthly_breakdown": False,
                },
                "spreadsheet": {
                    "file": "budgets.csv",
                    "key": "account",
                    "only": True,
                    "columns": {
                        "account_id": "Account ID",
                        "months": "2026-04",
                        "alerts": [],
                    },
                },
            }
        )
        config["spreadsheet"]["rows"] = {
            "111111111111": {
                "account_id": "111111111111",
                "amount": 47700,
                "monthly_amounts": [4000] * 9 + [3900] * 3,
                "account_name": "Platform",
                "folder": "Platform",
                "alerts": [],
                "warnings": [],
            }
        }

        budget = budget_for(
            config, "platform", "111111111111", "111111111111"
        )

        self.assertEqual(budget["period"], "YEARLY")
        self.assertTrue(budget["monthly_breakdown"])
        self.assertEqual(budget["start_time"], resolve_start("2026-04-01"))
        self.assertEqual(budget["amount"], 47700)
        self.assertEqual(budget["monthly_amounts"][-3:], [3900] * 3)

    def test_only_true_skips_accounts_without_a_row(self):
        config = _config(only=True)
        matched = budget_for(
            config,
            "hosted",
            "111111111111",
            "Alpha (111111111111)",
            now=datetime(2026, 1, 15, tzinfo=timezone.utc),
        )
        self.assertAlmostEqual(sum(matched["monthly_amounts"]), 400)
        self.assertEqual(matched["alerts"][0]["based_on"], "ACTUAL_COST")
        self.assertIsNone(budget_for(config, "hosted", "333333333333", "Gamma"))

    def test_only_false_keeps_the_yaml_amount(self):
        config = _config(only=False)
        kept = budget_for(config, "hosted", "333333333333", "Gamma (333333333333)")
        self.assertEqual(kept["amount"], 10)
        self.assertEqual(kept["alerts"], [])

    def test_account_name_mismatch_warns_and_still_matches(self):
        config = _config(only=True)
        budget = budget_for(
            config, "hosted", "111111111111", "Wrong (111111111111)"
        )
        self.assertIn("does not match Wrong", budget["_sheet_warnings"][0])
        self.assertEqual(budget["amount"], 400)

    def test_bucket_row_applies_to_every_account(self):
        config = _config(only=True, key="bucket")
        config["spreadsheet"]["rows"] = {
            "hosted": {
                "amount": 50,
                "alerts": [],
                "account_name": None,
                "warnings": [],
            }
        }
        first = budget_for(config, "hosted", "111111111111", "Alpha")
        second = budget_for(config, "hosted", "222222222222", "Beta")
        self.assertEqual(first["amount"], 50)
        self.assertEqual(second["amount"], 50)
        self.assertIsNone(budget_for(config, "other", "333333333333", "Gamma"))


def _config(only=True, key="account"):
    config = normalize_config(
        {
            "cost_category": "Example",
            "budget": {
                "type": "SPECIFIED_AMOUNT",
                "period": "YEARLY",
                "amount": 10,
                "monthly_breakdown": True,
                "start": "year_start",
            },
            "spreadsheet": {
                "file": "budgets.csv",
                "key": key,
                "only": only,
                "columns": {
                    "account_id": "Account ID",
                    "account_name": "Account Name",
                    "bucket": "Cost Bucket",
                    "amount": "Annual Budget",
                    "recipients": "Recipients",
                    "alerts": [
                        {
                            "header": "Alert 80",
                            "percentage": 80,
                            "based_on": "ACTUAL_COST",
                        }
                    ],
                },
            },
        }
    )
    config["spreadsheet"]["rows"] = {
        "111111111111": {
            "account_id": "111111111111",
            "amount": 400,
            "account_name": "Alpha",
            "alerts": [
                {
                    "percentage": 80,
                    "based_on": "ACTUAL_COST",
                    "emails": ["a@example.com"],
                }
            ],
            "warnings": [],
        }
    }
    return config


if __name__ == "__main__":
    unittest.main()
