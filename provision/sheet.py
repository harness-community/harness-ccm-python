from csv import DictReader
from datetime import date
from pathlib import Path


_YES = {"y", "yes", "true"}
_NO = {"n", "no", "false"}
_BASED_ON = {"ACTUAL_COST", "FORECASTED_COST"}
_KEYS = {"account", "bucket"}


def attach_spreadsheet(config: dict, config_path: str, override: str = None) -> dict:
    sheet = config.get("spreadsheet")
    if not sheet and not override:
        return config
    if not sheet:
        raise ValueError("spreadsheet column mapping is required to read a file")
    path = _sheet_path(config_path, sheet.get("file"), override)
    skipped = []
    loaded = load_spreadsheet(path, sheet, skipped)
    for message in skipped:
        print(message)
    sheet["file"] = str(path)
    sheet["rows"] = loaded
    return config


def load_spreadsheet(path, sheet: dict, skipped: list = None) -> dict:
    source = Path(path)
    if not source.is_file():
        raise ValueError(f"spreadsheet not found: {source}")
    columns = sheet["columns"]
    headers, records = _table(source, sheet.get("sheet"))
    _require_present(headers, columns, sheet["key"], source)
    key_header = columns["bucket"] if sheet["key"] == "bucket" else columns["account_id"]
    rows = {}
    for number, record in records:
        if _blank_row([record.get(key_header)]):
            if skipped is not None:
                skipped.append(
                    f"skipped spreadsheet row {number}: no {key_header}"
                )
            continue
        key = _row_key(record, columns, sheet["key"], number)
        if key in rows:
            label = "account id" if sheet["key"] == "account" else "bucket"
            raise ValueError(f"duplicate {label} {key} in {source}")
        rows[key] = _row(record, columns, sheet["key"], number)
    return rows


def unused_rows(sheet: dict, accounts: list) -> list:
    if not sheet or sheet.get("rows") is None:
        return []
    if sheet["key"] == "bucket":
        known = {account.bucket for account in accounts}
        noun = "bucket"
    else:
        known = {account.account_id for account in accounts}
        noun = "account"
    messages = []
    for key in sheet["rows"]:
        if key not in known:
            messages.append(
                f"skipped spreadsheet {noun} {key}: not in the selected cost category"
            )
    return messages


def _sheet_path(config_path: str, file_value: str, override: str = None) -> Path:
    chosen = override or file_value
    path = Path(chosen).expanduser()
    if path.is_absolute():
        return path
    if override:
        return Path.cwd() / path
    return Path(config_path).resolve().parent / path


def _table(path: Path, sheet_name: str):
    suffix = path.suffix.lower()
    if suffix == ".csv":
        if sheet_name:
            raise ValueError("spreadsheet.sheet applies to xlsx workbooks")
        return _csv_table(path)
    if suffix in (".xlsx", ".xlsm"):
        return _xlsx_table(path, sheet_name)
    raise ValueError(f"spreadsheet must be .xlsx or .csv: {path}")


def _csv_table(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = DictReader(handle)
        headers = _headers(reader.fieldnames, path)
        records = []
        for offset, row in enumerate(reader, start=2):
            values = [row.get(field) for field in reader.fieldnames or []]
            if _blank_row(values):
                continue
            records.append((offset, _record(headers, values)))
        return headers, records


def _xlsx_table(path: Path, sheet_name: str):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name and sheet_name not in workbook.sheetnames:
            raise ValueError(f"spreadsheet sheet not found: {sheet_name}")
        worksheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        pending = worksheet.iter_rows(values_only=True)
        try:
            header_row = next(pending)
        except StopIteration:
            raise ValueError(f"spreadsheet is empty: {path}") from None
        headers = _headers(header_row, path)
        records = []
        for offset, values in enumerate(pending, start=2):
            if _blank_row(values):
                continue
            records.append((offset, _record(headers, values)))
        return headers, records
    finally:
        workbook.close()


def _record(headers: list, values) -> dict:
    cells = list(values)
    record = {}
    for index, name in enumerate(headers):
        if not name or name in record:
            continue
        if headers.count(name) > 1:
            continue
        record[name] = cells[index] if index < len(cells) else None
    return record


def _require_present(headers: list, columns: dict, key: str, path: Path) -> None:
    required = []
    if columns.get("amount"):
        required.append(columns["amount"])
    if key == "account":
        required.append(columns["account_id"])
    else:
        required.append(columns["bucket"])
    if columns.get("account_name"):
        required.append(columns["account_name"])
    if columns.get("folder"):
        required.append(columns["folder"])
    if columns.get("recipients"):
        required.append(columns["recipients"])
    required.extend(columns.get("months") or [])
    required.extend(alert["header"] for alert in columns.get("alerts") or [])
    for header in required:
        positions = [index for index, name in enumerate(headers) if name == header]
        if not positions:
            raise ValueError(f"spreadsheet {path} is missing column {header}")
        if len(positions) > 1:
            letters = ", ".join(_column_letter(index) for index in positions)
            raise ValueError(
                f"spreadsheet {path} has more than one {header} column "
                f"({letters}); rename or remove the extra one"
            )


def _column_letter(index: int) -> str:
    letters = ""
    number = index + 1
    while number:
        number, remainder = divmod(number - 1, 26)
        letters = chr(ord("A") + remainder) + letters
    return letters


def _headers(headers, path: Path) -> list:
    if not headers:
        raise ValueError(f"spreadsheet is missing a header row: {path}")
    names = ["" if header is None else str(header).strip() for header in headers]
    if not any(names):
        raise ValueError(f"spreadsheet is missing a header row: {path}")
    return names


def _blank_row(values) -> bool:
    return all(value is None or str(value).strip() == "" for value in values)


def _row_key(record: dict, columns: dict, key: str, number: int) -> str:
    if key == "bucket":
        return _text(record, columns["bucket"], number, "bucket")
    return _account_id(record.get(columns["account_id"]), number)


def _row(record: dict, columns: dict, key: str, number: int) -> dict:
    monthly_amounts = [
        _amount(record.get(header), number) for header in columns.get("months") or []
    ]
    amount = (
        _amount(record.get(columns["amount"]), number)
        if columns.get("amount")
        else None
    )
    if monthly_amounts:
        total = sum(monthly_amounts)
        if amount is not None and abs(total - amount) > 0.02:
            warnings_for_total = (
                f"{columns['amount']} {amount:g} does not match the monthly columns {total:g}"
            )
        else:
            warnings_for_total = None
        amount = total
    else:
        warnings_for_total = None
    recipients = _recipients(record.get(columns["recipients"])) if columns.get("recipients") else []
    alerts = []
    warnings = []
    if warnings_for_total:
        warnings.append(warnings_for_total)
    for alert in columns.get("alerts") or []:
        header = alert["header"]
        if header not in record:
            raise ValueError(f"row {number} is missing column {header}")
        if not _enabled(record.get(header), number, header):
            continue
        if not recipients:
            warnings.append(f"{header} is enabled with no recipients")
        alerts.append(
            {
                "percentage": alert["percentage"],
                "based_on": alert["based_on"],
                "emails": list(recipients),
            }
        )
    name = None
    if columns.get("account_name"):
        raw_name = record.get(columns["account_name"])
        if raw_name is not None and str(raw_name).strip() != "":
            name = str(raw_name).strip()
    folder = None
    if columns.get("folder"):
        raw_folder = record.get(columns["folder"])
        if raw_folder is not None and str(raw_folder).strip() != "":
            folder = str(raw_folder).strip()
        else:
            warnings.append(
                f"{columns['folder']} is blank; using the cost bucket as the folder"
            )
    row = {
        "amount": amount,
        "monthly_amounts": monthly_amounts,
        "alerts": alerts,
        "account_name": name,
        "folder": folder,
        "warnings": warnings,
    }
    if key == "account":
        row["account_id"] = _account_id(record.get(columns["account_id"]), number)
    return row


def _account_id(value, number: int) -> str:
    if isinstance(value, bool) or value is None:
        raise ValueError(f"row {number} is missing an account id")
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not value.is_integer():
            raise ValueError(f"row {number} account id {value!r} is not a whole number")
        number_text = str(int(value))
        if int(value) < 0:
            raise ValueError(f"row {number} account id {value!r} is not a whole number")
        if len(number_text) <= 12:
            return number_text.zfill(12)
        return number_text
    text = str(value).strip()
    if not text:
        raise ValueError(f"row {number} is missing an account id")
    return text


def _amount(value, number: int) -> float:
    if isinstance(value, bool) or value is None or str(value).strip() == "":
        raise ValueError(f"row {number} is missing a budget amount")
    text = "".join(str(value).replace(",", "").replace("$", "").split())
    try:
        amount = float(text)
    except ValueError as exc:
        raise ValueError(f"row {number} budget amount {value!r} is not a number") from exc
    if amount < 0:
        raise ValueError(f"row {number} budget amount must be zero or greater")
    return amount


def _recipients(value) -> list:
    if value is None or str(value).strip() == "":
        return []
    return [part.strip() for part in str(value).split(",") if part.strip()]


def _enabled(value, number: int, header: str) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    text = str(value).strip().lower()
    if text == "" or text in _NO:
        return False
    if text in _YES:
        return True
    raise ValueError(f"row {number} column {header} must be Y or N")


def _text(record: dict, header: str, number: int, label: str) -> str:
    if header not in record:
        raise ValueError(f"row {number} is missing column {header}")
    value = record.get(header)
    text = "" if value is None else str(value).strip()
    if not text:
        raise ValueError(f"row {number} is missing a {label}")
    return text


def normalize_spreadsheet(raw) -> dict:
    if not raw:
        return None
    if not isinstance(raw, dict):
        raise ValueError("spreadsheet must be a mapping")
    key = raw.get("key", "account")
    if key not in _KEYS:
        raise ValueError("spreadsheet.key must be account or bucket")
    only = raw.get("only", False)
    if only not in (True, False):
        raise ValueError("spreadsheet.only must be true or false")
    file_value = raw.get("file")
    if not isinstance(file_value, str) or not file_value.strip():
        raise ValueError("spreadsheet.file is required")
    columns = raw.get("columns")
    if not isinstance(columns, dict):
        raise ValueError("spreadsheet.columns is required")
    months, month_start = _month_columns(columns.get("months"))
    normalized_columns = {
        "account_id": _column(columns, "account_id", key == "account"),
        "account_name": _column(columns, "account_name", False),
        "bucket": _column(columns, "bucket", key == "bucket"),
        "folder": _column(columns, "folder", False),
        "amount": _column(columns, "amount", not months),
        "recipients": _column(columns, "recipients", False),
        "months": months,
        "month_start": month_start,
        "alerts": _alert_columns(columns.get("alerts") or []),
    }
    return {
        "file": file_value,
        "sheet": raw.get("sheet") or None,
        "key": key,
        "only": only,
        "columns": normalized_columns,
        "rows": None,
    }


def _month_columns(months) -> tuple:
    if not months:
        return [], None
    if isinstance(months, str):
        try:
            start = date.fromisoformat(f"{months.strip()}-01")
        except ValueError as exc:
            raise ValueError(
                "spreadsheet.columns.months must be a YYYY-MM start header "
                "or a list of 12 column headers"
            ) from exc
        headers = []
        year = start.year
        month = start.month
        for _ in range(12):
            headers.append(f"{year:04d}-{month:02d}")
            month += 1
            if month > 12:
                month = 1
                year += 1
        return headers, start.isoformat()
    if not isinstance(months, list) or len(months) != 12:
        raise ValueError(
            "spreadsheet.columns.months must be a YYYY-MM start header "
            "or contain 12 column headers"
        )
    headers = []
    for index, header in enumerate(months):
        if not isinstance(header, str) or not header.strip():
            raise ValueError(f"spreadsheet.columns.months[{index}] must be a column header")
        headers.append(header)
    return headers, None


def _column(columns: dict, name: str, required: bool):
    value = columns.get(name)
    if value is None:
        if required:
            raise ValueError(f"spreadsheet.columns.{name} is required")
        return None
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"spreadsheet.columns.{name} must be a column header")
    return value


def _alert_columns(alerts) -> list:
    if not isinstance(alerts, list):
        raise ValueError("spreadsheet.columns.alerts must be a list")
    normalized = []
    for index, alert in enumerate(alerts):
        label = f"spreadsheet.columns.alerts[{index}]"
        if not isinstance(alert, dict):
            raise ValueError(f"{label} must be a mapping")
        header = alert.get("header")
        if not isinstance(header, str) or not header.strip():
            raise ValueError(f"{label}.header is required")
        percentage = alert.get("percentage")
        if isinstance(percentage, bool) or not isinstance(percentage, (int, float)):
            raise ValueError(f"{label}.percentage must be a number")
        based_on = alert.get("based_on", alert.get("basedOn", "ACTUAL_COST"))
        if based_on not in _BASED_ON:
            raise ValueError(f"{label}.based_on must be ACTUAL_COST or FORECASTED_COST")
        normalized.append(
            {"header": header, "percentage": percentage, "based_on": based_on}
        )
    return normalized
